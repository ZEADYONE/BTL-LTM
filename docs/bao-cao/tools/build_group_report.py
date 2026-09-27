from __future__ import annotations

import math
import os
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


from report_common import *  # noqa: F401,F403 - shared helpers, styles and paths

import chapter_3_ket_noi_xac_thuc
import chapter_4_sanh_phong_cho
import chapter_5_game_engine
import chapter_6_client_du_lieu_tran


def diagram_architecture() -> Path:
    image, draw = new_canvas(
        "KIẾN TRÚC TỔNG THỂ",
        "Client hiển thị và gửi thao tác; Server quản lý trạng thái; CSDL lưu dữ liệu lâu dài",
    )
    player = (70, 350, 320, 650)
    client = (405, 225, 780, 775)
    server = (900, 225, 1395, 775)
    db = (1480, 350, 1730, 650)
    for box in (player, client, server, db):
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=20, width=4)
    text_centered(draw, (95, 370, 315, 640), "NGƯỜI CHƠI\n\nNhập lệnh và theo dõi trận đấu", font(28, True), rgb(BLACK))
    text_centered(draw, (425, 245, 760, 305), "DESKTOP CLIENT", font(30, True), rgb(BLACK))
    for index, label in enumerate(["Giao diện", "Điều khiển", "Kết nối TCP", "Hiển thị trạng thái"]):
        y = 335 + index * 95
        box = (465, y, 720, y + 62)
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=12, width=3)
        text_centered(draw, box, label, font(23), rgb(BLACK))
    text_centered(draw, (930, 245, 1365, 305), "SERVER", font(30, True), rgb(BLACK))
    server_labels = [
        "Kết nối và phiên",
        "Tài khoản và sảnh",
        "Phòng chơi",
        "Bộ xử lý trò chơi",
        "Lịch sử và xếp hạng",
    ]
    for index, label in enumerate(server_labels):
        y = 315 + index * 88
        box = (965, y, 1330, y + 60)
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=12, width=3)
        text_centered(draw, box, label, font(22), rgb(BLACK))
    text_centered(draw, db, "CƠ SỞ DỮ LIỆU\n\nTài khoản\nKết quả trận\nBảng xếp hạng", font(25, True), rgb(BLACK))
    arrow(draw, (320, 500), (405, 500), BLACK)
    arrow(draw, (780, 455), (900, 455), BLACK)
    arrow(draw, (900, 555), (780, 555), BLACK)
    arrow(draw, (1395, 500), (1480, 500), BLACK)
    text_centered(draw, (790, 395, 890, 435), "Yêu cầu", font(21, True), rgb(BLACK))
    text_centered(draw, (790, 575, 890, 615), "Trạng thái", font(21), rgb(BLACK))
    return save_image(image, "hinh-2-1-kien-truc-tong-the.png")


def diagram_network() -> Path:
    image, draw = new_canvas(
        "GIAO TIẾP CLIENT - SERVER",
        "Mỗi Client duy trì một kết nối TCP; Server phân loại yêu cầu và trả dữ liệu phù hợp",
    )
    clients = [(80, 230, 370, 380), (80, 445, 370, 595), (80, 660, 370, 810)]
    for idx, box in enumerate(clients, 1):
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=18, width=4)
        text_centered(draw, box, f"CLIENT {idx}\nGửi thao tác - nhận trạng thái", font(23, True), rgb(BLACK))
    connection = (520, 270, 850, 770)
    processing = (1000, 270, 1370, 770)
    rooms = (1510, 330, 1730, 710)
    for box in (connection, processing, rooms):
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=20, width=4)
    text_centered(draw, connection, "QUẢN LÝ KẾT NỐI\n\nNhận dữ liệu\nXác định người dùng\nChuyển yêu cầu", font(26, True), rgb(BLACK))
    text_centered(draw, processing, "XỬ LÝ TRUNG TÂM\n\nĐăng nhập\nSảnh và phòng\nTrò chơi\nLịch sử - xếp hạng", font(26, True), rgb(BLACK))
    text_centered(draw, rooms, "NHIỀU PHÒNG\n\nPhòng A\nPhòng B\nPhòng C", font(25, True), rgb(BLACK))
    connection_points = [335, 520, 705]
    for box, target_y in zip(clients, connection_points):
        mid_y = (box[1] + box[3]) // 2
        arrow(draw, (box[2], mid_y - 14), (connection[0], target_y - 14), BLACK, width=5, head=16)
        arrow(draw, (connection[0], target_y + 22), (box[2], mid_y + 22), BLACK, width=5, head=16)
    arrow(draw, (connection[2], 455), (processing[0], 455), BLACK, width=6)
    arrow(draw, (processing[0], 585), (connection[2], 585), BLACK, width=6)
    arrow(draw, (processing[2], 455), (rooms[0], 455), BLACK, width=6)
    arrow(draw, (rooms[0], 585), (processing[2], 585), BLACK, width=6)
    text_centered(draw, (420, 835, 1380, 885), "Khung dữ liệu: độ dài + nội dung JSON", font(24), rgb(BLACK))
    return save_image(image, "hinh-2-2-giao-tiep-client-server.png")


def diagram_data_flow() -> Path:
    image, draw = new_canvas(
        "LUỒNG DỮ LIỆU TỔNG QUÁT",
        "Trạng thái hiển thị ở Client luôn được cập nhật từ kết quả xử lý của Server",
    )
    labels = [
        ("1", "Người chơi\nthao tác"),
        ("2", "Client tạo\nyêu cầu"),
        ("3", "Gửi qua\nTCP"),
        ("4", "Server kiểm tra\nvà xử lý"),
        ("5", "Cập nhật\ntrạng thái"),
        ("6", "Gửi kết quả\ncho Client"),
        ("7", "Client cập nhật\ngiao diện"),
    ]
    boxes = []
    x = 70
    for _, _ in labels:
        boxes.append((x, 345, x + 205, 635))
        x += 245
    for idx, (number, label) in enumerate(labels):
        box = boxes[idx]
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=18, width=4)
        draw.ellipse(
            (box[0] + 69, box[1] + 25, box[0] + 136, box[1] + 92),
            fill=rgb(WHITE),
            outline=rgb(BLACK),
            width=4,
        )
        num_width = draw.textbbox((0, 0), number, font=font(30, True))[2]
        draw.text((box[0] + 102 - num_width / 2, box[1] + 40), number, font=font(30, True), fill=rgb(BLACK))
        text_centered(draw, (box[0] + 10, box[1] + 105, box[2] - 10, box[3] - 15), label, font(24, True), rgb(BLACK))
        if idx < len(boxes) - 1:
            arrow(draw, (box[2] + 8, 490), (boxes[idx + 1][0] - 8, 490), BLACK, width=5, head=16)
    text_centered(
        draw,
        (300, 725, 1500, 790),
        "Client không tự quyết định vị trí, bom hoặc kết quả trận đấu.",
        font(28, True),
        rgb(BLACK),
    )
    return save_image(image, "hinh-2-3-luong-du-lieu.png")


def diagram_classes() -> Path:
    image, draw = new_canvas(
        "CÁC NHÓM LỚP CHÍNH",
        "Sơ đồ chỉ giữ các lớp đại diện để thể hiện trách nhiệm và quan hệ tổng thể",
    )
    client_group = (65, 210, 610, 825)
    server_group = (680, 210, 1735, 825)
    round_box(draw, client_group, rgb(WHITE), rgb(BLACK), radius=20, width=4)
    round_box(draw, server_group, rgb(WHITE), rgb(BLACK), radius=20, width=4)
    text_centered(draw, (140, 225, 535, 280), "PHÍA CLIENT", font(32, True), rgb(BLACK))
    text_centered(draw, (930, 225, 1490, 280), "PHÍA SERVER", font(32, True), rgb(BLACK))
    client_boxes = [
        ((120, 330, 555, 435), "Giao diện người dùng"),
        ((120, 500, 555, 605), "Điều khiển và trạng thái Client"),
        ((120, 670, 555, 775), "Kết nối mạng"),
    ]
    for box, label in client_boxes:
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=14, width=3)
        text_centered(draw, box, label, font(24, True), rgb(BLACK))
    server_boxes = [
        ((735, 315, 1045, 430), "Kết nối và phiên"),
        ((1125, 315, 1435, 430), "Xác thực và sảnh"),
        ((1515, 315, 1685, 430), "CSDL"),
        ((735, 530, 1045, 645), "Quản lý phòng"),
        ((1125, 530, 1435, 645), "Quản lý trận"),
        ((1515, 530, 1685, 645), "Kết quả"),
        ((930, 700, 1435, 790), "Trò chơi: bản đồ, người chơi, bom, vụ nổ"),
    ]
    for box, label in server_boxes:
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=14, width=3)
        text_centered(draw, box, label, font(22, True), rgb(BLACK))
    arrow(draw, (337, 435), (337, 500), BLACK, width=5)
    arrow(draw, (337, 605), (337, 670), BLACK, width=5)
    orthogonal_arrow(draw, [(555, 722), (645, 722), (645, 372), (735, 372)], BLACK, width=5)
    arrow(draw, (1045, 372), (1125, 372), BLACK, width=5)
    arrow(draw, (1435, 372), (1515, 372), BLACK, width=5)
    arrow(draw, (890, 430), (890, 530), BLACK, width=5)
    arrow(draw, (1045, 588), (1125, 588), BLACK, width=5)
    arrow(draw, (1280, 645), (1280, 700), BLACK, width=5)
    arrow(draw, (1435, 588), (1515, 588), BLACK, width=5)
    return save_image(image, "hinh-2-4-thiet-ke-lop.png")


def diagram_erd() -> Path:
    image, draw = new_canvas(
        "THIẾT KẾ DỮ LIỆU LƯU TRỮ",
        "Cơ sở dữ liệu chỉ lưu tài khoản và kết quả đã hoàn thành; trạng thái đang chơi được giữ trong bộ nhớ",
    )
    account = (70, 260, 520, 770)
    participant = (675, 260, 1125, 770)
    match = (1280, 260, 1730, 770)
    for box in (account, participant, match):
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=18, width=4)
        draw.line((box[0], 350, box[2], 350), fill=rgb(BLACK), width=4)
    text_centered(draw, (70, 275, 520, 335), "TÀI KHOẢN", font(28, True), rgb(BLACK))
    text_centered(draw, (675, 275, 1125, 335), "NGƯỜI THAM GIA", font(28, True), rgb(BLACK))
    text_centered(draw, (1280, 275, 1730, 335), "TRẬN ĐẤU", font(28, True), rgb(BLACK))
    text_centered(draw, (100, 385, 490, 735), "Mã tài khoản\nTên đăng nhập\nMật khẩu đã mã hóa\nTổng điểm\nSố trận thắng - thua - hòa", font(24), rgb(BLACK))
    text_centered(draw, (705, 385, 1095, 735), "Mã trận\nMã tài khoản\nTên người chơi\nKết quả cá nhân\nĐiểm nhận được", font(24), rgb(BLACK))
    text_centered(draw, (1310, 385, 1700, 735), "Mã trận\nMã phòng\nThời điểm bắt đầu - kết thúc\nNgười thắng\nKết quả chung", font(24), rgb(BLACK))
    arrow(draw, (520, 485), (675, 485), BLACK, width=6)
    arrow(draw, (1280, 610), (1125, 610), BLACK, width=6)
    text_centered(draw, (540, 425, 655, 470), "1        N", font(24, True), rgb(BLACK))
    text_centered(draw, (1145, 550, 1260, 595), "N        1", font(24, True), rgb(BLACK))
    text_centered(
        draw,
        (560, 835, 1240, 890),
        "Mỗi dòng người tham gia liên kết một tài khoản với một trận đấu.",
        font(23),
        rgb(BLACK),
    )
    return save_image(image, "hinh-2-5-erd.png")


def diagram_game_flow() -> Path:
    image, draw = new_canvas(
        "LUỒNG XỬ LÝ MỘT TRẬN ĐẤU",
        "Các thao tác được kiểm tra tại Server trước khi tạo trạng thái mới cho cả phòng",
    )
    steps = [
        "Phòng chờ",
        "Người chơi\nsẵn sàng",
        "Server khởi tạo\ntrận đấu",
        "Nhận và kiểm tra\nthao tác",
        "Cập nhật\ntrạng thái",
        "Gửi trạng thái\ncho cả phòng",
        "Kết thúc và\nlưu kết quả",
    ]
    positions = [
        (90, 260, 380, 400),
        (510, 260, 800, 400),
        (930, 260, 1220, 400),
        (1350, 260, 1640, 400),
        (1350, 625, 1640, 765),
        (930, 625, 1220, 765),
        (510, 625, 800, 765),
    ]
    for label, box in zip(steps, positions):
        round_box(draw, box, rgb(WHITE), rgb(BLACK), radius=18, width=4)
        text_centered(draw, box, label, font(24, True), rgb(BLACK))
    for first, second in zip(positions[:4], positions[1:4]):
        arrow(draw, (first[2] + 10, 330), (second[0] - 10, 330), BLACK, width=6)
    arrow(draw, (1455, 400), (1455, 625), BLACK, width=6)
    arrow(draw, (1350, 695), (1220, 695), BLACK, width=6)
    arrow(draw, (930, 695), (800, 695), BLACK, width=6)
    orthogonal_arrow(draw, [(1640, 695), (1700, 695), (1700, 330), (1648, 330)], BLACK, width=5)
    text_centered(draw, (1515, 485, 1680, 565), "Lặp lại đến khi\ncó kết quả", font(21, True), rgb(BLACK))
    return save_image(image, "hinh-2-6-luong-tran-dau.png")


ASSIGNMENTS = [
    # (STT, module, nội dung lập trình mạng, chương)
    (
        "1",
        "Kết nối và xác thực",
        "Server TCP, phiên kết nối, giao thức đóng khung và mã hóa thông điệp, tầng kết nối phía Client, đăng ký, đăng nhập, trạng thái trực tuyến, xử lý mất kết nối.",
        "Chương 3",
    ),
    (
        "2",
        "Sảnh và phòng chờ",
        "Đồng bộ danh sách người trực tuyến và danh sách phòng; tạo, tham gia, rời phòng; sẵn sàng, bắt đầu, chơi lại; phát trạng thái phòng tới các Client liên quan.",
        "Chương 4",
    ),
    (
        "3",
        "Game Engine và đồng bộ thời gian thực",
        "Nhận thao tác di chuyển, đặt bom; vòng lặp 20 tick/giây cho từng phòng; phát trạng thái trận và kết quả; thu phím và hiển thị trận đấu ở Client.",
        "Chương 5",
    ),
    (
        "4",
        "Kiến trúc Desktop Client và dữ liệu trận đấu",
        "Tầng ứng dụng Client: gửi lệnh, phân phối thông điệp, cập nhật giao diện; lưu kết quả trận, tính điểm; truy vấn lịch sử và bảng xếp hạng.",
        "Chương 6",
    ),
]


def add_assignment_page(doc):
    """Assignment table right after the cover; names are filled in by the group later."""
    add_heading(doc, "BẢNG PHÂN CÔNG NHIỆM VỤ", "CTDT-H0")
    add_body(
        doc,
        "Phần chung gồm Mở đầu, Chương 1, Chương 2 và Kết luận do cả nhóm thực hiện. Nội dung cá nhân của từng thành viên được trình bày trong chương tương ứng theo bảng dưới đây.",
    )
    add_table(
        doc,
        None,
        ["STT", "Họ và tên", "Mã sinh viên", "Module phụ trách", "Nội dung lập trình mạng", "Chương"],
        [[stt, "", "", module, content, chapter] for stt, module, content, chapter in ASSIGNMENTS],
        [600, 1700, 1300, 1700, 3000, 1000],
    )


def add_contents_lists(doc):
    """MỤC LỤC and the lists of figures and tables as Word fields (refreshed with F9)."""
    for heading, instruction in (
        ("MỤC LỤC", 'TOC \\o "1-3" \\h \\z \\u'),
        ("DANH MỤC CÁC HÌNH VẼ", 'TOC \\h \\z \\t "CTDT-Hinh,1"'),
        ("DANH MỤC CÁC BẢNG BIỂU", 'TOC \\h \\z \\t "CTDT-Bang,1"'),
    ):
        add_heading(doc, heading, "CTDT-H0")
        paragraph = doc.add_paragraph(style="CTDT-Text")
        paragraph.paragraph_format.first_line_indent = Inches(0)
        add_field(paragraph, instruction)


def finalize_document(doc, output: Path):
    """A4 page size, page numbers from the first content section, metadata, black text, save."""
    for section in doc.sections:
        section.page_width = Inches(8.2677)
        section.page_height = Inches(11.6929)
    if len(doc.sections) >= 2:
        add_page_number(doc.sections[1])
    else:
        add_page_number(doc.sections[0])

    doc.core_properties.title = "Báo cáo bài tập lớn - Bomberman Online Mini"
    doc.core_properties.subject = "Kiến trúc, thiết kế và cài đặt hệ thống"
    doc.core_properties.author = "Nhóm sinh viên thực hiện"
    doc.core_properties.keywords = "Bomberman, lập trình mạng, TCP, Client-Server, Java"

    force_all_text_black(doc)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(output)


def build_report():
    diagrams = {
        "architecture": diagram_architecture(),
        "network": diagram_network(),
        "flow": diagram_data_flow(),
        "classes": diagram_classes(),
        "erd": diagram_erd(),
        "game_flow": diagram_game_flow(),
    }

    doc = Document(str(TEMPLATE))
    replace_cover_text(doc)
    clear_after_cover(doc)
    configure_styles(doc)
    enable_update_fields(doc)

    # Front matter. The retained section break at the end of the cover starts this content on a new page;
    # every CTDT-H0 heading inherits page-break-before from CTDT-H1.
    add_assignment_page(doc)
    add_contents_lists(doc)

    add_heading(doc, "MỞ ĐẦU", "CTDT-H0")
    add_body(
        doc,
        "Sự phát triển của các ứng dụng tương tác qua mạng đặt ra yêu cầu không chỉ truyền được dữ liệu mà còn phải duy trì trạng thái thống nhất giữa nhiều người dùng. Với trò chơi đối kháng trực tuyến, yêu cầu này thể hiện rõ ở việc mọi người chơi phải nhìn thấy cùng một bản đồ, vị trí, bom, vụ nổ và kết quả trong cùng một thời điểm hợp lý.",
    )
    add_body(
        doc,
        "Đề tài Bomberman Online Mini xây dựng một trò chơi desktop cho từ hai đến bốn người trong một phòng và cho phép nhiều phòng hoạt động đồng thời. Hệ thống áp dụng mô hình một Server phục vụ nhiều Client. Client tiếp nhận thao tác và hiển thị, trong khi Server kiểm tra yêu cầu, cập nhật trạng thái và gửi kết quả trở lại các máy đang tham gia.",
    )
    add_body(
        doc,
        "Báo cáo gồm hai phần. Phần chung, gồm Chương 1 và Chương 2, trình bày tổng quan, kiến trúc chung và thiết kế chung của hệ thống. Phần cá nhân, gồm Chương 3 đến Chương 6, trình bày chi tiết từng module theo bảng phân công: kết nối và xác thực, sảnh và phòng chờ, Game Engine và đồng bộ thời gian thực, kiến trúc Desktop Client và dữ liệu trận đấu. Mỗi chương cá nhân mô tả chức năng bằng biểu đồ, giao diện và có bảng ánh xạ giữa hình vẽ với chức năng đã thực hiện.",
    )

    # Chapter 1
    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN HỆ THỐNG", "CTDT-H1")
    add_heading(doc, "1.1 Giới thiệu bài toán", "CTDT-H2")
    add_body(
        doc,
        "Bomberman Online Mini là trò chơi đối kháng trực tuyến, trong đó người chơi di chuyển trên bản đồ dạng ô, đặt bom để phá vật cản và loại đối thủ. Khác với trò chơi chạy trên một máy, phiên bản trực tuyến phải tiếp nhận thao tác từ nhiều máy và bảo đảm tất cả người chơi cùng theo dõi một trạng thái thống nhất.",
    )
    add_body(
        doc,
        "Bài toán được giải quyết bằng một máy chủ trung tâm. Máy khách chỉ gửi ý định, chẳng hạn hướng di chuyển hoặc yêu cầu đặt bom. Máy chủ kiểm tra điều kiện hợp lệ, cập nhật trận đấu và gửi trạng thái mới cho các máy trong cùng phòng. Nhờ đó, người chơi không tự thay đổi vị trí hay kết quả ở phía máy của mình.",
    )

    add_heading(doc, "1.2 Mục tiêu của hệ thống", "CTDT-H2")
    add_body(doc, "Hệ thống được xây dựng hướng tới các mục tiêu chính sau:")
    for item in [
        "Cho phép nhiều người chơi kết nối, đăng ký, đăng nhập và sử dụng hệ thống qua mạng TCP.",
        "Tổ chức sảnh chờ và nhiều phòng chơi độc lập, mỗi phòng có từ hai đến bốn người.",
        "Duy trì một trạng thái trận đấu thống nhất do Server quản lý.",
        "Cập nhật kịp thời trạng thái phòng và trạng thái trò chơi tới đúng người chơi liên quan.",
        "Lưu kết quả đã hoàn thành để cung cấp lịch sử và bảng xếp hạng.",
        "Giữ kiến trúc đủ rõ để có thể thay đổi giao diện Client mà không làm thay đổi luật xử lý phía Server.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1.3 Phạm vi chức năng", "CTDT-H2")
    add_heading(doc, "1.3.1 Chức năng phía người chơi", "CTDT-H3")
    for item in [
        "Đăng ký tài khoản, đăng nhập và đăng xuất.",
        "Xem người chơi đang trực tuyến và danh sách phòng.",
        "Tạo phòng, tham gia phòng, rời phòng và thay đổi trạng thái sẵn sàng.",
        "Di chuyển nhân vật, đặt bom và theo dõi diễn biến trận đấu.",
        "Xem kết quả, lịch sử các trận đã chơi và bảng xếp hạng.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1.3.2 Chức năng phía Server", "CTDT-H3")
    for item in [
        "Quản lý kết nối, phiên đăng nhập và trạng thái trực tuyến.",
        "Quản lý danh sách phòng, thành viên, chủ phòng và trạng thái sẵn sàng.",
        "Kiểm tra thao tác của người chơi và điều khiển diễn biến trận đấu.",
        "Gửi trạng thái phòng, trạng thái trận đấu và kết quả tới đúng Client.",
        "Lưu tài khoản, kết quả trận và số liệu phục vụ xếp hạng.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1.4 Công nghệ sử dụng", "CTDT-H2")
    add_body(
        doc,
        "Các công nghệ được lựa chọn theo yêu cầu của một ứng dụng desktop có giao tiếp mạng liên tục. Bảng 1.1 trình bày vai trò của từng nhóm công nghệ ở mức tổng quát.",
    )
    add_table(
        doc,
        "Bảng 1.1. Công nghệ chính của hệ thống",
        ["Nhóm công nghệ", "Công nghệ", "Vai trò trong hệ thống"],
        [
            ["Ngôn ngữ và xây dựng", "Java 21, Gradle", "Phát triển và quản lý bốn module của dự án."],
            ["Server", "Spring Boot", "Khởi tạo ứng dụng Server và quản lý các thành phần xử lý."],
            ["Giao diện", "JavaFX", "Xây dựng giao diện desktop chính cho người chơi."],
            ["Giao tiếp mạng", "TCP Socket, JSON", "Duy trì kết nối và truyền thông điệp hai chiều."],
            ["Lưu trữ", "MySQL, JPA", "Lưu tài khoản, kết quả trận và số liệu xếp hạng."],
            ["Bảo mật mật khẩu", "BCrypt", "Lưu mật khẩu dưới dạng đã mã hóa một chiều."],
        ],
        [2250, 2150, 4900],
    )

    add_heading(doc, "1.5 Kết chương", "CTDT-H2")
    add_body(
        doc,
        "Chương 1 đã giới thiệu bài toán, mục tiêu, phạm vi chức năng và các công nghệ chính của Bomberman Online Mini. Trọng tâm của hệ thống là duy trì trạng thái chung cho nhiều người chơi thông qua một Server trung tâm. Chương tiếp theo trình bày cách tổ chức các thành phần và thiết kế tổng thể để thực hiện mục tiêu đó.",
    )

    # Chapter 2
    add_heading(doc, "CHƯƠNG 2. KIẾN TRÚC VÀ THIẾT KẾ HỆ THỐNG", "CTDT-H1")
    add_heading(doc, "2.1 Kiến trúc tổng thể", "CTDT-H2")
    add_body(
        doc,
        "Hệ thống được tổ chức theo mô hình Client-Server. Nhiều Desktop Client kết nối tới một Server thông qua TCP. Client chịu trách nhiệm giao tiếp với người dùng, thu nhận thao tác và trình bày dữ liệu. Server quản lý tài khoản, sảnh, phòng chơi, diễn biến trận đấu và dữ liệu lưu trữ.",
    )
    add_body(
        doc,
        "Trong thời gian thi đấu, Server là nơi giữ trạng thái chính. Mỗi Client chỉ gửi thao tác điều khiển; kết quả sau khi kiểm tra được Server gửi lại dưới dạng trạng thái hoàn chỉnh của trận. Cách tổ chức này giúp các máy không phải tự tính toán theo những bản dữ liệu riêng và hạn chế chênh lệch giữa người chơi.",
    )
    add_body(
        doc,
        "Dự án được chia thành bốn module. Module dùng chung chứa định dạng thông điệp; module Server chứa xử lý trung tâm; module JavaFX là Client chính; module libGDX là Client tương thích với cùng giao thức. Hình 2.1 thể hiện các khối chính và quan hệ giữa chúng.",
    )
    add_figure(
        doc,
        diagrams["architecture"],
        "Hình 2.1. Kiến trúc tổng thể của hệ thống",
        "Sơ đồ người chơi sử dụng Desktop Client, trao đổi TCP với Server và Server lưu dữ liệu vào cơ sở dữ liệu.",
    )

    add_heading(doc, "2.2 Các thành phần chính", "CTDT-H2")
    add_body(
        doc,
        "Mỗi thành phần có một trách nhiệm rõ ràng. Bảng 2.1 chỉ trình bày các khối có ảnh hưởng trực tiếp đến kiến trúc chung; các thành phần hiển thị nhỏ và chi tiết xử lý nội bộ không được liệt kê.",
    )
    add_table(
        doc,
        "Bảng 2.1. Các thành phần chính của hệ thống",
        ["Thành phần", "Trách nhiệm", "Dữ liệu nhận", "Kết quả tạo ra"],
        [
            ["Desktop Client", "Hiển thị giao diện và nhận thao tác người chơi.", "Thao tác người dùng; dữ liệu từ Server.", "Yêu cầu mạng; giao diện được cập nhật."],
            ["Kết nối TCP", "Duy trì kênh trao đổi hai chiều giữa Client và Server.", "Thông điệp cần gửi; dữ liệu từ socket.", "Thông điệp đã gửi hoặc đã nhận."],
            ["Quản lý phiên", "Theo dõi từng kết nối và người dùng đã đăng nhập.", "Kết nối mới; thông tin đăng nhập.", "Phiên hoạt động hoặc thông báo lỗi."],
            ["Sảnh và phòng", "Quản lý người online, danh sách phòng, thành viên và trạng thái sẵn sàng.", "Yêu cầu tạo, tham gia, rời phòng.", "Trạng thái sảnh và phòng mới."],
            ["Bộ xử lý trò chơi", "Kiểm tra thao tác, cập nhật bản đồ, người chơi, bom và kết quả.", "Lệnh của người chơi; thời gian xử lý.", "Trạng thái trận đấu mới."],
            ["Lưu trữ dữ liệu", "Lưu tài khoản và kết quả sau trận.", "Dữ liệu tài khoản; kết quả hoàn thành.", "Lịch sử và bảng xếp hạng."],
        ],
        [1900, 2850, 2200, 2350],
    )

    add_heading(doc, "2.3 Kiến trúc giao tiếp Client-Server", "CTDT-H2")
    add_body(
        doc,
        "Server mở cổng TCP 8081 theo cấu hình mặc định. Mỗi Client tạo một kết nối và duy trì kết nối này trong suốt phiên sử dụng. Nhờ kết nối liên tục, Server có thể chủ động gửi trạng thái mới mà không cần Client tạo lại kết nối cho từng thao tác.",
    )
    add_body(
        doc,
        "Mỗi thông điệp gồm phần cho biết độ dài và phần nội dung JSON. Phần độ dài giúp bên nhận xác định chính xác điểm kết thúc của một thông điệp trong dòng dữ liệu TCP. Nội dung JSON gồm loại thông điệp, mã yêu cầu và dữ liệu đi kèm. Mã yêu cầu giúp Client biết phản hồi tương ứng với thao tác nào.",
    )
    add_body(
        doc,
        "Khi có nhiều Client, Server duy trì riêng thông tin của từng kết nối. Dữ liệu phòng chỉ được gửi cho thành viên trong phòng; dữ liệu trận đấu chỉ được gửi cho người tham gia trận đó. Hình 2.2 mô tả cách nhiều Client cùng trao đổi với Server và cách Server tách xử lý theo từng phòng.",
    )
    add_figure(
        doc,
        diagrams["network"],
        "Hình 2.2. Kiến trúc giao tiếp giữa Client và Server",
        "Sơ đồ nhiều Client kết nối TCP tới Server, Server xử lý trung tâm và quản lý nhiều phòng độc lập.",
    )

    add_heading(doc, "2.4 Luồng dữ liệu tổng quát", "CTDT-H2")
    add_body(
        doc,
        "Luồng dữ liệu bắt đầu từ thao tác của người chơi. Client chuyển thao tác thành yêu cầu và gửi qua kết nối TCP. Server xác định người gửi, kiểm tra trạng thái phòng và điều kiện của trò chơi trước khi chấp nhận. Sau khi trạng thái được cập nhật, kết quả mới được gửi tới các Client liên quan để cập nhật giao diện.",
    )
    add_body(
        doc,
        "Client không gửi tọa độ đích và không tự thông báo rằng một người chơi đã thắng. Ví dụ, với thao tác di chuyển, Client chỉ gửi hướng. Server kiểm tra ô đích có nằm trong bản đồ, có bị tường, bom hoặc người chơi khác chiếm giữ hay không. Chỉ sau bước này vị trí mới được thay đổi.",
    )
    add_figure(
        doc,
        diagrams["flow"],
        "Hình 2.3. Luồng dữ liệu tổng quát trong hệ thống",
        "Chuỗi từ thao tác người chơi qua Client, TCP, Server, cập nhật trạng thái rồi quay lại giao diện Client.",
    )

    add_heading(doc, "2.5 Thiết kế module", "CTDT-H2")
    add_body(
        doc,
        "Việc chia module giúp tách phần giao tiếp chung khỏi giao diện và phần xử lý Server. Bảng 2.2 mô tả vai trò và quan hệ chính của từng module mà không đi sâu vào các tệp hoặc hàm cụ thể.",
    )
    add_table(
        doc,
        "Bảng 2.2. Thiết kế các module của dự án",
        ["Module", "Vai trò", "Quan hệ chính"],
        [
            ["common", "Định nghĩa thông điệp, dữ liệu trao đổi và cách đóng gói dữ liệu TCP.", "Được Server và cả hai Client sử dụng."],
            ["server", "Quản lý kết nối, tài khoản, phòng, trò chơi và lưu trữ.", "Nhận yêu cầu từ Client và kết nối tới cơ sở dữ liệu."],
            ["client-fx", "Giao diện desktop chính bằng JavaFX.", "Gửi thao tác tới Server và hiển thị trạng thái nhận được."],
            ["client", "Giao diện libGDX tương thích với giao thức hiện tại.", "Dùng chung module common và kết nối cùng Server."],
        ],
        [1800, 4100, 3400],
    )
    add_body(
        doc,
        "Bên trong module Server, mã nguồn tiếp tục được chia theo các nhóm mạng, xác thực, sảnh, phòng, trò chơi, kết quả và lưu trữ. Đây là cách phân lớp theo trách nhiệm trong cùng một ứng dụng, không phải kiến trúc nhiều dịch vụ độc lập.",
    )

    add_heading(doc, "2.6 Thiết kế lớp ở mức tổng quát", "CTDT-H2")
    add_body(
        doc,
        "Sơ đồ lớp trong phần nhóm chỉ lựa chọn các lớp đại diện. Mục tiêu là cho thấy đường đi từ giao diện tới kết nối, từ kết nối tới quản lý phòng và trò chơi, sau đó tới lưu trữ. Các lớp hỗ trợ hiển thị, hiệu ứng và kiểm tra nhỏ không được đưa vào để sơ đồ dễ đọc.",
    )
    add_body(
        doc,
        "Phía Client có ba nhóm chính: giao diện, điều khiển/trạng thái và kết nối mạng. Phía Server bắt đầu từ quản lý kết nối, sau đó chuyển yêu cầu tới nhóm xác thực, sảnh, phòng hoặc trò chơi. Khi trận kết thúc, nhóm kết quả chuyển dữ liệu cần lưu tới cơ sở dữ liệu.",
    )
    add_figure(
        doc,
        diagrams["classes"],
        "Hình 2.4. Thiết kế các nhóm lớp chính",
        "Sơ đồ nhóm lớp phía Client và Server, thể hiện đường đi từ giao diện tới trò chơi và cơ sở dữ liệu.",
    )

    add_heading(doc, "2.7 Thiết kế dữ liệu", "CTDT-H2")
    add_body(
        doc,
        "Dữ liệu của hệ thống được chia thành dữ liệu lưu lâu dài và dữ liệu đang hoạt động. Tài khoản, kết quả trận và số liệu xếp hạng được lưu trong MySQL. Danh sách người online, phòng đang có và trạng thái trận đang chạy được giữ trong bộ nhớ Server để đáp ứng nhanh.",
    )
    add_body(
        doc,
        "Thiết kế lưu trữ gồm ba nhóm dữ liệu chính. Bảng tài khoản lưu thông tin đăng nhập và thống kê. Bảng trận đấu lưu thời gian, phòng và kết quả chung. Bảng người tham gia lưu kết quả và điểm của từng người trong mỗi trận.",
    )
    add_table(
        doc,
        "Bảng 2.3. Dữ liệu chính được lưu trữ",
        ["Nhóm dữ liệu", "Thông tin chính", "Mục đích"],
        [
            ["Tài khoản", "Mã tài khoản, tên đăng nhập, mật khẩu đã mã hóa, tổng điểm, số trận thắng - thua - hòa.", "Xác thực người dùng và lập bảng xếp hạng."],
            ["Trận đấu", "Mã phòng, thời điểm bắt đầu - kết thúc, người thắng, kết quả chung.", "Lưu một lần thi đấu đã hoàn thành."],
            ["Người tham gia", "Tài khoản, tên người chơi, kết quả cá nhân và điểm nhận được.", "Thể hiện chi tiết từng người trong một trận."],
        ],
        [2100, 4300, 2900],
    )
    add_figure(
        doc,
        diagrams["erd"],
        "Hình 2.5. Mô hình dữ liệu lưu trữ",
        "Sơ đồ dữ liệu gồm tài khoản, trận đấu và người tham gia; một trận có nhiều người tham gia.",
    )
    add_body(
        doc,
        "Trạng thái đang chơi không được ghi liên tục xuống cơ sở dữ liệu. Khi trận kết thúc, Server mới lưu kết quả và cập nhật thống kê. Lựa chọn này làm giảm thao tác lưu trữ trong lúc trận đang diễn ra, nhưng trận chưa hoàn thành sẽ không được phục hồi nếu Server dừng đột ngột.",
    )

    add_heading(doc, "2.8 Thiết kế giao thức trao đổi dữ liệu", "CTDT-H2")
    add_body(
        doc,
        "Thông điệp được phân theo nhóm chức năng để hai phía cùng hiểu ý nghĩa dữ liệu. Báo cáo không liệt kê chi tiết mọi trường dữ liệu; Bảng 2.4 chỉ thể hiện các nhóm cần thiết cho kiến trúc chung.",
    )
    add_table(
        doc,
        "Bảng 2.4. Các nhóm thông điệp Client-Server",
        ["Nhóm", "Chiều chính", "Nội dung tiêu biểu", "Kết quả"],
        [
            ["Kết nối", "Hai chiều", "Kiểm tra kết nối và thông báo lỗi.", "Xác định kết nối còn hoạt động hoặc có lỗi."],
            ["Tài khoản", "Client yêu cầu, Server trả lời", "Đăng ký, đăng nhập, đăng xuất.", "Tạo hoặc xác nhận phiên người dùng."],
            ["Sảnh", "Hai chiều", "Danh sách người online và danh sách phòng.", "Cập nhật màn hình sảnh."],
            ["Phòng", "Hai chiều", "Tạo, tham gia, rời phòng, sẵn sàng, bắt đầu.", "Cập nhật trạng thái phòng cho thành viên."],
            ["Điều khiển trận", "Client gửi", "Hướng di chuyển và yêu cầu đặt bom.", "Được kiểm tra trước khi cập nhật."],
            ["Trạng thái trận", "Server gửi", "Bản đồ, người chơi, bom, vụ nổ và kết quả.", "Các Client trong phòng hiển thị cùng dữ liệu."],
            ["Sau trận", "Hai chiều", "Chơi lại, lịch sử và bảng xếp hạng.", "Chuẩn bị vòng mới hoặc hiển thị dữ liệu đã lưu."],
        ],
        [1550, 2200, 3100, 2450],
    )
    add_body(
        doc,
        "Trong source có khai báo một loại thông điệp riêng cho sự kiện người chơi bị loại, nhưng luồng hiện tại không gửi loại thông điệp này. Trạng thái sống hoặc bị loại được truyền trong trạng thái chung của trận. Vì vậy, báo cáo không xem thông điệp riêng này là một chức năng đang hoạt động.",
    )

    add_heading(doc, "2.9 Các luồng xử lý chính", "CTDT-H2")
    add_heading(doc, "2.9.1 Đăng nhập", "CTDT-H3")
    add_body(
        doc,
        "Client gửi tên đăng nhập và mật khẩu tới Server. Server kiểm tra dữ liệu tài khoản, đối chiếu mật khẩu đã mã hóa và xác định tài khoản có đang online ở nơi khác hay không. Nếu hợp lệ, người dùng được gắn với kết nối hiện tại và Client chuyển tới màn hình chính.",
    )

    add_heading(doc, "2.9.2 Tạo hoặc tham gia phòng", "CTDT-H3")
    add_body(
        doc,
        "Khi tạo phòng, người dùng trở thành chủ phòng. Khi tham gia, Server kiểm tra phòng còn tồn tại, đang chờ và chưa đủ bốn người. Sau mỗi thay đổi, trạng thái phòng đầy đủ được gửi lại cho các thành viên, nhờ đó danh sách người chơi và trạng thái sẵn sàng được cập nhật đồng nhất.",
    )

    add_heading(doc, "2.9.3 Bắt đầu và vận hành trận đấu", "CTDT-H3")
    add_body(
        doc,
        "Chủ phòng chỉ có thể bắt đầu khi phòng có ít nhất hai người và tất cả thành viên đã sẵn sàng. Server tạo trạng thái ban đầu, đặt người chơi tại các vị trí xuất phát và mở bộ xử lý riêng cho phòng. Các thao tác sau đó được tiếp nhận, kiểm tra và xử lý theo thứ tự.",
    )
    add_body(
        doc,
        "Trận đấu được cập nhật theo nhịp đều đặn. Trạng thái mới được gửi nhiều lần trong một giây để Client hiển thị chuyển động, bom và vụ nổ. Mỗi phòng có bộ xử lý riêng, vì vậy một trận ở phòng này không sử dụng trạng thái của phòng khác.",
    )
    add_figure(
        doc,
        diagrams["game_flow"],
        "Hình 2.6. Luồng xử lý một trận đấu",
        "Luồng từ phòng chờ, sẵn sàng, khởi tạo trận, xử lý thao tác, gửi trạng thái đến kết thúc và lưu kết quả.",
    )

    add_heading(doc, "2.9.4 Kết thúc trận và mất kết nối", "CTDT-H3")
    add_body(
        doc,
        "Khi chỉ còn một người sống, người đó thắng; nếu mọi người bị loại cùng lúc, trận hòa. Server gửi kết quả cho các Client, chuyển phòng sang trạng thái đã kết thúc và lưu kết quả. Điểm được tính theo đơn vị nguyên để biểu diễn thắng một điểm, hòa nửa điểm và thua không điểm.",
    )
    add_body(
        doc,
        "Nếu một người mất kết nối trong lúc chơi, Server chuyển sự kiện này vào cùng luồng xử lý của trận thay vì sửa trạng thái ngay tại phần kết nối. Người chơi bị đánh dấu đã bị loại, sau đó kết quả được tính theo các quy tắc chung. Cách xử lý này giữ thứ tự và tránh hai nơi cùng thay đổi trạng thái trận.",
    )

    add_heading(doc, "2.10 Đồng bộ và hỗ trợ nhiều phòng", "CTDT-H2")
    add_body(
        doc,
        "Mỗi phòng đang thi đấu có vùng trạng thái và bộ xử lý riêng. Các thao tác của phòng được đưa tới đúng bộ xử lý, còn trạng thái chỉ được gửi tới các kết nối thuộc phòng đó. Nhờ sự tách biệt này, Server có thể vận hành nhiều phòng đồng thời mà không trộn người chơi hoặc dữ liệu giữa các trận.",
    )
    add_body(
        doc,
        "Việc gửi trạng thái hoàn chỉnh thay vì chỉ gửi phần thay đổi làm thông điệp lớn hơn, nhưng giúp Client dễ khôi phục nếu bỏ lỡ một lần cập nhật. Mỗi bản nhận mới có thể thay thế trạng thái hiển thị trước đó, giảm phụ thuộc vào lịch sử thông điệp ở phía Client.",
    )

    add_heading(doc, "2.11 Một số giới hạn của thiết kế hiện tại", "CTDT-H2")
    for item in [
        "Kết nối TCP hiện chưa có lớp mã hóa, do đó chỉ phù hợp cho môi trường học tập hoặc mạng tin cậy nếu chưa bổ sung bảo mật đường truyền.",
        "Trạng thái trận đang diễn ra được giữ trong bộ nhớ và không được phục hồi sau khi Server khởi động lại.",
        "Client chưa có cơ chế tiếp tục phiên thi đấu sau khi mất kết nối.",
        "Cấu hình cơ sở dữ liệu cần được thống nhất với hướng dẫn chạy để tránh đặt thông tin nhạy cảm trực tiếp trong source.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.12 Kết chương", "CTDT-H2")
    add_body(
        doc,
        "Chương 2 đã trình bày kiến trúc Client-Server, các thành phần chính, cách trao đổi dữ liệu, thiết kế module, lớp và dữ liệu của Bomberman Online Mini. Thiết kế đặt Server ở vị trí quản lý trạng thái trung tâm, đồng thời tách riêng xử lý của từng phòng. Đây là cơ sở để nhiều Client cùng tham gia nhưng vẫn quan sát một diễn biến thống nhất. Các chương tiếp theo trình bày chi tiết cài đặt của từng module trên nền kiến trúc này.",
    )

    # Individual chapters, one per member (see ASSIGNMENTS).
    missing_screenshots: list[str] = []
    for chapter in (
        chapter_3_ket_noi_xac_thuc,
        chapter_4_sanh_phong_cho,
        chapter_5_game_engine,
        chapter_6_client_du_lieu_tran,
    ):
        missing_screenshots += chapter.write(doc).missing_screenshots

    add_heading(doc, "KẾT LUẬN", "CTDT-H0")
    add_body(
        doc,
        "Báo cáo đã trình bày hệ thống Bomberman Online Mini từ kiến trúc chung tới cài đặt của từng module. Hệ thống đáp ứng mô hình một Server - nhiều Client qua kết nối TCP duy trì liên tục, với giao thức đóng khung bằng tiền tố độ dài và nội dung JSON. Server quản lý phiên, tài khoản, sảnh và phòng trong bộ nhớ, xử lý trận đấu theo mô hình có thẩm quyền với một vòng lặp riêng cho mỗi phòng, lưu kết quả và số liệu xếp hạng vào cơ sở dữ liệu. Client JavaFX gửi thao tác, nhận trạng thái và hiển thị mà không tự quyết định diễn biến trận.",
    )
    add_body(
        doc,
        "Việc chia module theo tính năng giúp mỗi phần có trách nhiệm rõ ràng: kết nối và xác thực cung cấp phiên TCP đã định danh, sảnh và phòng chờ đồng bộ trạng thái chung trước trận, Game Engine đồng bộ trạng thái trận theo thời gian thực, còn Desktop Client và dữ liệu trận đấu kết nối các phần đó với người dùng và lưu lại kết quả. Các điểm giao giữa module như bắt đầu trận, kết thúc trận và mất kết nối đều đi qua những giao diện lập trình xác định.",
    )
    add_body(
        doc,
        "Hệ thống vẫn còn một số giới hạn đã nêu ở mục 2.11, như chưa mã hóa đường truyền, chưa khôi phục trận sau khi Server khởi động lại và chưa hỗ trợ tiếp tục phiên sau khi mất kết nối. Hướng phát triển tiếp theo gồm bổ sung TLS cho kết nối TCP, cơ chế kết nối lại vào trận đang diễn ra, giới hạn tần suất thao tác phía Server và quản lý phiên bản lược đồ cơ sở dữ liệu.",
    )

    add_heading(doc, "TÀI LIỆU THAM KHẢO", "CTDT-H0")
    references = [
        "[1] Nhóm phát triển, Mã nguồn dự án Bomberman Online Mini, phiên bản tại thời điểm lập báo cáo, 2026.",
        "[2] Tệp README và cấu hình xây dựng của dự án Bomberman Online Mini, 2026.",
        "[3] Tài liệu kỹ thuật đi kèm mã nguồn Java 21, Spring Boot, JavaFX và MySQL được sử dụng trong dự án.",
    ]
    for ref in references:
        add_reference(doc, ref)

    finalize_document(doc, OUTPUT)
    if missing_screenshots:
        print("Ảnh giao diện còn thiếu (đang dùng khung giữ chỗ), lưu vào", SCREENSHOT_DIR)
        for name in missing_screenshots:
            print("  -", name)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Dựng báo cáo Word của nhóm từ template.")
    parser.add_argument("--output", type=Path, help="Đường dẫn .docx khác, ví dụ khi file chính đang mở trong Word")
    args = parser.parse_args()
    if args.output is not None:
        OUTPUT = args.output.resolve()
    build_report()
