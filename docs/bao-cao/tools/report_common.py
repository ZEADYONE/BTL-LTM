"""Shared drawing and python-docx helpers for the Bomberman report builders."""
from __future__ import annotations

import math
import os
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor



ROOT = Path(__file__).resolve().parents[3]
TEMPLATE = ROOT / "docs" / "bao-cao" / "doc temp.docx"
OUTPUT = ROOT / "docs" / "bao-cao" / "bao_cao_nhom_bomberman.docx"
ASSET_DIR = ROOT / "docs" / "bao-cao" / "generated-report-assets"
# Real UI captures for the individual chapters; missing files are replaced by a labelled placeholder.
SCREENSHOT_DIR = ROOT / "docs" / "bao-cao" / "screenshots"
# Images the group inserted by hand in Word (extracted by docx_to_python.py).
IMAGE_DIR = ROOT / "docs" / "bao-cao" / "images"

BLUE = "000000"
BLUE_LIGHT = "000000"
BLUE_PALE = "FFFFFF"
GREEN = "000000"
GREEN_LIGHT = "FFFFFF"
ORANGE = "000000"
ORANGE_LIGHT = "FFFFFF"
GRAY = "000000"
GRAY_LIGHT = "FFFFFF"
RED = "000000"
WHITE = "FFFFFF"
BLACK = "000000"



def rgb(hex_value: str) -> tuple[int, int, int]:
    return tuple(int(hex_value[i : i + 2], 16) for i in (0, 2, 4))


FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def wrapped_lines(draw: ImageDraw.ImageDraw, text: str, chosen_font, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines():
        if not raw_line.strip():
            lines.append("")
            continue
        words = raw_line.split()
        current = ""
        for word in words:
            candidate = word if not current else current + " " + word
            if draw.textbbox((0, 0), candidate, font=chosen_font)[2] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines or [""]


def text_centered(draw, box, text, chosen_font, fill, spacing=8):
    x1, y1, x2, y2 = box
    lines = wrapped_lines(draw, text, chosen_font, max(10, x2 - x1 - 34))
    heights = [draw.textbbox((0, 0), line, font=chosen_font)[3] for line in lines]
    total_height = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_height) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=chosen_font)[2]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=chosen_font, fill=fill)
        y += height + spacing


def round_box(draw, box, fill, outline, radius=24, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color=BLACK, width=7, head=20):
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill=rgb(color), width=width)
    angle = math.atan2(ey - sy, ex - sx)
    left = (
        ex - head * math.cos(angle) + head * 0.55 * math.sin(angle),
        ey - head * math.sin(angle) - head * 0.55 * math.cos(angle),
    )
    right = (
        ex - head * math.cos(angle) - head * 0.55 * math.sin(angle),
        ey - head * math.sin(angle) + head * 0.55 * math.cos(angle),
    )
    draw.polygon([(ex, ey), left, right], fill=rgb(color))


def orthogonal_arrow(draw, points, color=BLACK, width=6, head=18):
    """Draw an aligned connector and place one arrow head on its final segment."""
    for start, end in zip(points, points[1:-1]):
        draw.line((*start, *end), fill=rgb(color), width=width)
    arrow(draw, points[-2], points[-1], color=color, width=width, head=head)


def new_canvas(title: str, subtitle: str | None = None):
    image = Image.new("RGB", (1800, 1000), rgb(WHITE))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((25, 25, 1775, 975), radius=26, outline=rgb(BLACK), width=3)
    draw.text((70, 55), title, font=font(42, True), fill=rgb(BLACK))
    if subtitle:
        draw.text((70, 115), subtitle, font=font(24), fill=rgb(BLACK))
    draw.line((70, 160, 1730, 160), fill=rgb(BLACK), width=3)
    return image, draw


def save_image(image: Image.Image, name: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / name
    image.save(path, format="PNG", optimize=True)
    return path


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color=BLACK, size=8):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths: list[int], indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, name="Times New Roman", size=13, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)


def set_style_font(style, size, bold=None, italic=None, color=BLACK):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "Times New Roman")


def set_outline_level(style, level: int):
    p_pr = style.element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def configure_styles(doc: Document):
    style_specs = {
        "CTDT-H0": (14, True, False, BLACK),
        "CTDT-H1": (14, True, False, BLACK),
        "CTDT-H2": (14, True, False, BLACK),
        "CTDT-H3": (13, True, True, BLACK),
        "CTDT-H4": (13, False, True, BLACK),
        "CTDT-Text": (13, False, False, BLACK),
        "CTDT-Bullet1": (13, False, False, BLACK),
        "CTDT-Bullet2": (13, False, False, BLACK),
        "CTDT-Hinh": (12, False, True, BLACK),
        "CTDT-Bang": (12, False, True, BLACK),
        "CTDT-NumerRef": (12, False, False, BLACK),
    }
    for name, (size, bold, italic, color) in style_specs.items():
        set_style_font(doc.styles[name], size, bold, italic, color)
    for name, level in (("CTDT-H0", 0), ("CTDT-H1", 0), ("CTDT-H2", 1), ("CTDT-H3", 2), ("CTDT-H4", 3)):
        set_outline_level(doc.styles[name], level)

    doc.styles["CTDT-H0"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["CTDT-H0"].paragraph_format.space_before = Pt(12)
    doc.styles["CTDT-H0"].paragraph_format.space_after = Pt(12)
    doc.styles["CTDT-H0"].paragraph_format.keep_with_next = True

    doc.styles["CTDT-H1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["CTDT-H1"].paragraph_format.space_before = Pt(0)
    doc.styles["CTDT-H1"].paragraph_format.space_after = Pt(18)
    doc.styles["CTDT-H1"].paragraph_format.keep_with_next = True
    doc.styles["CTDT-H1"].paragraph_format.page_break_before = True

    doc.styles["CTDT-H2"].paragraph_format.space_before = Pt(12)
    doc.styles["CTDT-H2"].paragraph_format.space_after = Pt(6)
    doc.styles["CTDT-H2"].paragraph_format.keep_with_next = True
    doc.styles["CTDT-H2"].paragraph_format.page_break_before = False

    doc.styles["CTDT-H3"].paragraph_format.space_before = Pt(9)
    doc.styles["CTDT-H3"].paragraph_format.space_after = Pt(4)
    doc.styles["CTDT-H3"].paragraph_format.keep_with_next = True
    doc.styles["CTDT-H3"].paragraph_format.page_break_before = False
    doc.styles["CTDT-H4"].paragraph_format.page_break_before = False

    body = doc.styles["CTDT-Text"].paragraph_format
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.first_line_indent = Inches(0.39)
    body.space_before = Pt(0)
    body.space_after = Pt(6)
    body.line_spacing = 1.3

    for name, left, hanging in (("CTDT-Bullet1", 0.39, 0.24), ("CTDT-Bullet2", 0.72, 0.24)):
        pf = doc.styles[name].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Inches(left)
        pf.first_line_indent = Inches(-hanging)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.2

    for name in ("CTDT-Hinh", "CTDT-Bang"):
        pf = doc.styles[name].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = name == "CTDT-Bang"
        pf.space_before = Pt(6)
        pf.space_after = Pt(9)


def clear_after_cover(doc: Document):
    body = doc._element.body
    start = None
    for index, child in enumerate(list(body)):
        if child.tag == qn("w:p"):
            texts = child.findall(".//" + qn("w:t"))
            value = "".join(node.text or "" for node in texts).strip()
            if value == "MỤC LỤC":
                start = index
                break
    if start is None:
        raise RuntimeError("Không tìm thấy điểm bắt đầu MỤC LỤC trong template")
    for child in list(body)[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def replace_cover_text(doc: Document):
    replacements = {
        "HỌC PHẦN: <TÊN HỌC PHẦN>": "HỌC PHẦN: LẬP TRÌNH MẠNG",
        "MÃ HỌC PHẦN: <MÃ HỌC PHẦN>": "MÃ HỌC PHẦN: [CẦN BỔ SUNG]",
        "ĐỀ TÀI: <TÊN ĐỀ TÀI>": "ĐỀ TÀI: BOMBERMAN ONLINE MINI - TRÒ CHƠI ĐỐI KHÁNG TRỰC TUYẾN",
        "<Mã sinh viên><Họ tên sinh viên>": "[CẦN BỔ SUNG MÃ SINH VIÊN VÀ HỌ TÊN]",
        "Tên nhóm: <số nhóm>": "Tên nhóm: [CẦN BỔ SUNG]",
        "Tên lớp: <Tên lớp>": "Tên lớp: [CẦN BỔ SUNG]",
        "Giảng viên hướng dẫn: <Chức danh> + <Họ tên GV>": "Giảng viên hướng dẫn: [CẦN BỔ SUNG]",
        "HÀ NỘI 2024": "HÀ NỘI 2026",
    }
    for paragraph in doc.paragraphs:
        old = paragraph.text.strip()
        if old not in replacements:
            continue
        text = replacements[old]
        for run in paragraph.runs:
            run.text = ""
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        run.text = text
        size = 13
        bold = old.startswith(("HỌC PHẦN:", "MÃ HỌC PHẦN:", "ĐỀ TÀI:"))
        if old.startswith("ĐỀ TÀI:"):
            size = 14
        if old.startswith("HÀ NỘI"):
            bold = True
        set_run_font(run, size=size, bold=bold)


def add_field(paragraph, instruction: str, placeholder: str = "Cập nhật trường trong Microsoft Word (F9)"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=12, color=BLACK)


def enable_update_fields(doc: Document):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_heading(doc, text: str, style="CTDT-H2"):
    # The supplied template already numbers chapter and section heading styles.
    # Keep source labels readable while avoiding duplicated numbers in Word.
    if style == "CTDT-H1":
        text = re.sub(r"^CHƯƠNG\s+\d+\.\s*", "", text, flags=re.IGNORECASE)
    elif style in {"CTDT-H2", "CTDT-H3"}:
        text = re.sub(r"^\d+(?:\.\d+)+\s*", "", text)
    return doc.add_paragraph(text, style=style)


def disable_style_numbering(paragraph):
    """Keep explicit chapter-based figure/table labels instead of style numbering."""
    properties = paragraph._p.get_or_add_pPr()
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    numbering = OxmlElement("w:numPr")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "0")
    numbering.append(number_id)
    properties.append(numbering)


def add_body(doc, text: str, bold_prefix: str | None = None, bold: bool = False):
    paragraph = doc.add_paragraph(style="CTDT-Text")
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=13, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest, size=13)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=13, bold=True if bold else None)
    return paragraph


def add_blank(doc):
    """Empty body paragraph used as extra vertical space (e.g. before a figure)."""
    return doc.add_paragraph(style="CTDT-Text")


def add_reference(doc, text: str):
    """One entry of TÀI LIỆU THAM KHẢO; the CTDT-NumerRef style numbers it."""
    paragraph = doc.add_paragraph(style="CTDT-NumerRef")
    run = paragraph.add_run(re.sub(r"^\[\d+\]\s*", "", text))
    set_run_font(run, size=12)
    return paragraph


def add_bullet(doc, text: str, level=1):
    # Use a visible Unicode bullet so the list renders consistently across Word versions.
    paragraph = doc.add_paragraph(style="CTDT-Text")
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.left_indent = Inches(0.45 if level == 1 else 0.78)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(f"• {text}")
    set_run_font(run, size=13)
    return paragraph


def add_figure(doc, path: Path, caption: str, alt: str, width: float = 6.10):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    drawing = run._element.find(".//" + qn("wp:docPr"))
    if drawing is not None:
        drawing.set("title", caption)
        drawing.set("descr", alt)
    caption_paragraph = doc.add_paragraph(style="CTDT-Hinh")
    disable_style_numbering(caption_paragraph)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=12, italic=True)


def add_table(doc, caption: str | None, headers: list[str], rows: list[list[str]], widths: list[int]):
    # caption=None keeps the table out of the list of tables (e.g. the assignment page).
    if caption is not None:
        caption_paragraph = doc.add_paragraph(style="CTDT-Bang")
        disable_style_numbering(caption_paragraph)
        caption_run = caption_paragraph.add_run(caption)
        set_run_font(caption_run, size=12, italic=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table, color=BLACK, size=8)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, WHITE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, size=11, bold=True, color=BLACK)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_shading(cell, WHITE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            set_run_font(run, size=10.5)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_page_number(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=11)
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), "1")


def force_all_text_black(doc: Document):
    """Remove residual theme/accent font colors from the template and content."""
    for style in doc.styles:
        try:
            style.font.color.rgb = RGBColor.from_string(BLACK)
        except (AttributeError, ValueError):
            pass

    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(BLACK)
            color = run._element.get_or_add_rPr().find(qn("w:color"))
            if color is not None:
                color.attrib.pop(qn("w:themeColor"), None)
                color.attrib.pop(qn("w:themeTint"), None)
                color.attrib.pop(qn("w:themeShade"), None)
                color.set(qn("w:val"), BLACK)



def add_code(doc, code: str):
    """Short source excerpt in a bordered single-cell table (not listed in the list of tables)."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_borders(table, color=BLACK, size=6)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, WHITE)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    lines = code.strip("\n").splitlines()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=9.5)
        if index < len(lines) - 1:
            run.add_break()
    set_table_geometry(table, [9300])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def screenshot_placeholder(file_name: str, hint: str) -> Path:
    """Draws a clearly marked frame standing in for a UI capture that has not been taken yet."""
    image = Image.new("RGB", (1600, 900), rgb(WHITE))
    draw = ImageDraw.Draw(image)
    x1, y1, x2, y2 = 30, 30, 1570, 870
    dash, gap = 26, 14
    for x in range(x1, x2, dash + gap):
        draw.line((x, y1, min(x + dash, x2), y1), fill=rgb(BLACK), width=4)
        draw.line((x, y2, min(x + dash, x2), y2), fill=rgb(BLACK), width=4)
    for y in range(y1, y2, dash + gap):
        draw.line((x1, y, x1, min(y + dash, y2)), fill=rgb(BLACK), width=4)
        draw.line((x2, y, x2, min(y + dash, y2)), fill=rgb(BLACK), width=4)
    text_centered(draw, (120, 190, 1480, 300), "VỊ TRÍ CHÈN ẢNH GIAO DIỆN", font(46, True), rgb(BLACK))
    text_centered(draw, (160, 320, 1440, 620), hint, font(32), rgb(BLACK), spacing=12)
    text_centered(
        draw,
        (160, 660, 1440, 780),
        f"Chụp từ ứng dụng đang chạy và lưu tại\ndocs/bao-cao/screenshots/{file_name}",
        font(28, True),
        rgb(BLACK),
    )
    return save_image(image, "placeholder-" + Path(file_name).stem + ".png")


def add_screenshot(doc, file_name: str, caption: str, alt: str, hint: str) -> bool:
    """Inserts screenshots/<file_name> when present; otherwise a placeholder. Returns True if real."""
    path = SCREENSHOT_DIR / file_name
    real = path.is_file()
    add_figure(doc, path if real else screenshot_placeholder(file_name, hint), caption, alt)
    return real


class FigureLog:
    """Numbers the figures of one chapter and collects the figure ↔ function mapping table."""

    def __init__(self, chapter: int):
        self.chapter = chapter
        self.count = 0
        self.rows: list[list[str]] = []
        self.missing_screenshots: list[str] = []

    def _label(self, number: str) -> str:
        self.count += 1
        expected = f"{self.chapter}.{self.count}"
        if number != expected:
            raise ValueError(f"Figure numbered {number}, expected {expected}")
        return f"Hình {number}"

    def figure(self, doc, number: str, path: Path, title: str, alt: str, function: str, source: str):
        label = self._label(number)
        add_figure(doc, path, f"{label}. {title}", alt)
        self.rows.append([label, title, function, source])

    def screenshot(self, doc, number: str, file_name: str, title: str, alt: str, hint: str, function: str, source: str):
        label = self._label(number)
        if not add_screenshot(doc, file_name, f"{label}. {title}", alt, hint):
            self.missing_screenshots.append(file_name)
        self.rows.append([label, title, function, source])

    def mapping_table(self, doc, caption: str):
        add_table(
            doc,
            caption,
            ["Mã hình", "Hình vẽ / giao diện", "Chức năng được minh họa", "Thành phần cài đặt"],
            self.rows,
            [1100, 2600, 3000, 2600],
        )
