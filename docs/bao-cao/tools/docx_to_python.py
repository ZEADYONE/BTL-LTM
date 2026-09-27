"""Sinh các module nội dung Python (tools/final/*.py) từ một báo cáo Word đã chỉnh tay.

Dùng khi nhóm sửa trực tiếp file .docx: chạy lại script này để phần Python khớp với bản Word,
sau đó dựng lại bằng build_final_report.py.

    python docs/bao-cao/tools/docx_to_python.py [--input <file.docx>] [--out-dir <thư mục>]

Chỉ đọc nội dung sau trang bìa. Mục lục và danh mục hình/bảng là trường Word nên được dựng lại
bằng build_final_report.py. Ảnh trùng với ảnh trong generated-report-assets/ hoặc screenshots/
được tham chiếu tới file đó; ảnh chèn tay trong Word được tách ra docs/bao-cao/images/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from report_common import ASSET_DIR, IMAGE_DIR, ROOT, SCREENSHOT_DIR

DEFAULT_INPUT = ROOT / "docs" / "bao-cao" / "bao_cao_nhom_bomberman-final.docx"
DEFAULT_OUT_DIR = Path(__file__).resolve().parent / "final"

EMU_PER_INCH = 914400
DEFAULT_FIGURE_WIDTH = 6.10
BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
DOC_PR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"

# Named image folders, in the order they are searched for a matching picture.
IMAGE_SOURCES = [("ASSET_DIR", ASSET_DIR), ("SCREENSHOT_DIR", SCREENSHOT_DIR), ("IMAGE_DIR", IMAGE_DIR)]

CHAPTER_MODULES = {
    3: ("chuong_3_ket_noi_xac_thuc", "Chương 3 - Kết nối và xác thực (phần cá nhân A)"),
    4: ("chuong_4_sanh_phong_cho", "Chương 4 - Sảnh và phòng chờ (phần cá nhân B)"),
    5: ("chuong_5_game_engine", "Chương 5 - Game Engine và đồng bộ thời gian thực (phần cá nhân C)"),
    6: ("chuong_6_client_du_lieu_tran", "Chương 6 - Kiến trúc Desktop Client và dữ liệu trận đấu (phần cá nhân D)"),
}
CONTENTS_HEADINGS = {"MỤC LỤC", "DANH MỤC CÁC HÌNH VẼ", "DANH MỤC CÁC BẢNG BIỂU"}


def py(value: str) -> str:
    """Double-quoted Python string literal (JSON string syntax is valid Python)."""
    return json.dumps(value, ensure_ascii=False)


def slug(text: str) -> str:
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


class Module:
    def __init__(self, name: str, title: str):
        self.name = name
        self.title = title
        self.lines: list[str] = []
        self.imports: set[str] = set()

    def call(self, function: str, *args: str, multiline: bool = False):
        self.imports.add(function)
        if not multiline:
            self.lines.append(f"    {function}(doc, {', '.join(args)})")
            return
        self.lines.append(f"    {function}(")
        self.lines.append("        doc,")
        for arg in args:
            self.lines.append(f"        {arg},")
        self.lines.append("    )")

    def source(self, input_name: str) -> str:
        names = sorted(self.imports)
        header = [
            f'"""{self.title}.',
            "",
            f"Sinh bởi docx_to_python.py từ {input_name}, sau đó có thể sửa tay.",
            "Chạy lại docx_to_python.py sẽ ghi đè file này.",
            '"""',
            "from __future__ import annotations",
            "",
        ]
        if names:
            header += ["from report_common import (", *[f"    {name}," for name in names], ")"]
        return "\n".join(header + ["", "", "def write(doc):", *self.lines]) + "\n"


class Converter:
    def __init__(self, source: Path):
        self.source = source
        self.document = docx.Document(str(source))
        self.known_images = self._index_images()
        self.modules: list[Module] = []
        self.front: list[str] = []
        self.body: list[str] = []
        self.chapter = 0
        self.section = 0
        self.subsection = 0
        self.warnings: list[str] = []

    @staticmethod
    def _index_images() -> dict[str, tuple[str, str]]:
        known = {}
        for variable, folder in IMAGE_SOURCES:
            if folder.is_dir():
                for path in sorted(folder.iterdir()):
                    if path.is_file() and not path.name.startswith("placeholder-"):
                        known.setdefault(hashlib.md5(path.read_bytes()).hexdigest(), (variable, path.name))
        return known

    def _start_module(self, name: str, title: str, front: bool):
        module = Module(name, title)
        self.modules.append(module)
        (self.front if front else self.body).append(name)
        return module

    def convert(self):
        module: Module | None = None
        skipping_contents = False
        pending_image = None
        pending_table_caption = None
        previous_was_table = False
        started = False

        for element in self.document.element.body.iterchildren():
            if element.tag == qn("w:sectPr"):
                continue
            if element.tag == qn("w:tbl"):
                if not started or skipping_contents:
                    continue
                self._table(module, Table(element, self.document), pending_table_caption)
                pending_table_caption = None
                previous_was_table = True
                continue

            paragraph = Paragraph(element, self.document)
            style = paragraph.style.name
            text = paragraph.text.strip()

            if style == "CTDT-H0":
                if text == "BẢNG PHÂN CÔNG NHIỆM VỤ":
                    started = True
                    module = self._start_module("phan_cong", "Bảng phân công nhiệm vụ (ngay sau trang bìa)", True)
                elif text in CONTENTS_HEADINGS:
                    skipping_contents = True
                    continue
                elif text == "MỞ ĐẦU":
                    skipping_contents = False
                    module = self._start_module("mo_dau_chuong_1_2", "Mở đầu, Chương 1 và Chương 2 (phần nhóm)", False)
                elif text == "KẾT LUẬN":
                    module = self._start_module("ket_luan", "Kết luận và tài liệu tham khảo", False)
            if not started or skipping_contents:
                continue

            if style == "CTDT-H1":
                self.chapter += 1
                self.section = self.subsection = 0
                if self.chapter in CHAPTER_MODULES:
                    module = self._start_module(*CHAPTER_MODULES[self.chapter], False)
                elif self.chapter > 2:
                    module = self._start_module(f"chuong_{self.chapter}", f"Chương {self.chapter}", False)

            if element.findall(".//" + BLIP):
                pending_image = self._image(element)
                previous_was_table = False
                continue
            if style == "CTDT-Hinh":
                if pending_image is None:
                    raise ValueError(f"Chú thích hình không có ảnh đi kèm: {text}")
                self._figure(module, pending_image, text)
                pending_image = None
                continue
            if pending_image is not None:
                raise ValueError(f"Ảnh không có chú thích CTDT-Hinh ngay sau (trước: {text[:60]})")

            if style == "CTDT-Bang":
                pending_table_caption = text
                continue
            if not text and style == "Normal" and previous_was_table:
                previous_was_table = False  # the spacer add_table() inserts after every table
                continue
            previous_was_table = False
            self._paragraph(module, paragraph, style, text)

        return self

    def _paragraph(self, module: Module, paragraph: Paragraph, style: str, text: str):
        if style == "CTDT-H0":
            module.call("add_heading", py(text), py("CTDT-H0"))
        elif style == "CTDT-H1":
            module.call("add_heading", py(f"CHƯƠNG {self.chapter}. {text}"), py("CTDT-H1"))
        elif style == "CTDT-H2":
            self.section += 1
            self.subsection = 0
            module.call("add_heading", py(f"{self.chapter}.{self.section} {text}"), py("CTDT-H2"))
        elif style == "CTDT-H3":
            self.subsection += 1
            module.call("add_heading", py(f"{self.chapter}.{self.section}.{self.subsection} {text}"), py("CTDT-H3"))
        elif style == "CTDT-NumerRef":
            module.call("add_reference", py(text))
        elif style == "CTDT-Text":
            properties = paragraph._p.pPr
            is_list = properties is not None and properties.numPr is not None
            if not text:
                module.call("add_blank")
            elif text.startswith("• "):
                module.call("add_bullet", py(text[2:]))
            elif is_list:
                module.call("add_bullet", py(text))
            else:
                runs = [run for run in paragraph.runs if run.text.strip()]
                bold = bool(runs) and all(run.bold for run in runs)
                args = [py(text)] + (["bold=True"] if bold else [])
                module.call("add_body", *args, multiline=len(text) > 90)
        else:
            self.warnings.append(f"Bỏ qua đoạn style {style}: {text[:60]}")

    def _image(self, element):
        blip = element.findall(".//" + BLIP)[0]
        part = self.document.part.related_parts[blip.get(qn("r:embed"))]
        blob = part.blob
        extent = element.findall(".//" + EXTENT)[0]
        doc_pr = element.findall(".//" + DOC_PR)[0]
        return {
            "blob": blob,
            "extension": Path(str(part.partname)).suffix or ".png",
            "width": round(int(extent.get("cx")) / EMU_PER_INCH, 2),
            "alt": doc_pr.get("descr"),
        }

    def _figure(self, module: Module, image: dict, caption: str):
        digest = hashlib.md5(image["blob"]).hexdigest()
        if digest not in self.known_images:
            label = re.match(r"(Hình \d+\.\d+)\.?\s*(.*)", caption)
            name = slug(label[1] + " " + label[2]) if label else slug(caption)
            file_name = name + image["extension"]
            IMAGE_DIR.mkdir(parents=True, exist_ok=True)
            (IMAGE_DIR / file_name).write_bytes(image["blob"])
            self.known_images[digest] = ("IMAGE_DIR", file_name)
            print("Tách ảnh chèn tay:", IMAGE_DIR / file_name)
        variable, file_name = self.known_images[digest]
        module.imports.add(variable)
        alt = image["alt"] or re.sub(r"^Hình \d+\.\d+\.\s*", "", caption)
        args = [f"{variable} / {py(file_name)}", py(caption), py(alt)]
        if abs(image["width"] - DEFAULT_FIGURE_WIDTH) > 0.005:
            args.append(f"width={image['width']}")
        module.call("add_figure", *args, multiline=True)

    def _table(self, module: Module, table: Table, caption: str | None):
        grid = table._tbl.tblGrid
        widths = [int(column.get(qn("w:w"))) for column in grid.findall(qn("w:gridCol"))]
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        module.imports.add("add_table")
        module.lines += ["    add_table(", "        doc,", f"        {py(caption) if caption else 'None'},"]
        module.lines.append(f"        [{', '.join(py(value) for value in rows[0])}],")
        module.lines.append("        [")
        for row in rows[1:]:
            module.lines.append(f"            [{', '.join(py(value) for value in row)}],")
        module.lines += ["        ],", f"        {widths},", "    )"]

    def write(self, out_dir: Path):
        out_dir.mkdir(parents=True, exist_ok=True)
        for module in self.modules:
            (out_dir / f"{module.name}.py").write_text(module.source(self.source.name), encoding="utf-8")
            print("Đã sinh", out_dir / f"{module.name}.py")
        init = [
            f'"""Nội dung báo cáo sinh từ {self.source.name}; thứ tự dựng do build_final_report.py dùng."""',
            "",
            "# Trước MỤC LỤC (ngay sau trang bìa).",
            f"FRONT = {self.front!r}",
            "# Sau danh mục hình và bảng.",
            f"BODY = {self.body!r}",
        ]
        (out_dir / "__init__.py").write_text("\n".join(init) + "\n", encoding="utf-8")
        for warning in self.warnings:
            print("Cảnh báo:", warning)


def main():
    parser = argparse.ArgumentParser(description="Sinh module Python từ báo cáo Word đã chỉnh tay.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="File .docx nguồn")
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR, help="Thư mục ghi các module")
    args = parser.parse_args()
    Converter(args.input.resolve()).convert().write(args.out_dir.resolve())


if __name__ == "__main__":
    main()
